"""Build the requested plain workbook from Python output snapshots using openpyxl."""

import json
from math import isfinite
from datetime import date
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
import pandas as pd

from mini_tms import config as cfg
from mini_tms.config import C

# CONFIG: Python outputs are the workbook inputs; financial drivers are illustrative.
INPUT_DIRECTORY = cfg.OUTPUT
WORKBOOK_PATH = cfg.WORKBOOK
EXPECTED_PATH = WORKBOOK_PATH.with_name("reconciliation_expected.json")
USE_AVERAGE_COLUMN = "use_average_balances"
CAP_FRACTION_COLUMN = "regulatory_cap_fraction"
CARRY_COLUMN = "carry_excess_eca"
DECIMAL_FORMAT = '#,##0.00;(#,##0.00);0.00'
PERCENT_FORMAT = "0.00%"
NUMBER_WIDTH = 19
TEXT_WIDTH = 28
NO_FINITE_BALANCE = "No finite balance"


def build_workbook(input_directory=INPUT_DIRECTORY, path=WORKBOOK_PATH):
    """Write live formulas and separate Python expectations for engine reconciliation."""
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = cfg.AUTHOR
    book.properties.lastModifiedBy = cfg.AUTHOR
    book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True)
    expected = {}

    def sheet(name, records, columns, formulas, percentage_columns=()):
        ws = book.create_sheet(name)
        ws.append(["Amounts in USD; CCC ratios in days. Illustrative inputs unless provenance states SEC 10-K."])
        ws.append(columns)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        ws["A1"].font = Font(name="Arial", size=10)
        ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 30
        lookup = {key: get_column_letter(number) for number, key in enumerate(columns, start=1)}
        for number, record in enumerate(records, start=3):
            for key in columns:
                cell = ws[f"{lookup[key]}{number}"]
                value = record.get(key)
                if key in formulas:
                    cell.value = formulas[key](number, lookup)
                    result = float(record[key])
                    if not isfinite(result):
                        if key != C.REQUIRED_BALANCE or result != float("inf"):
                            raise ValueError(f"Nonfinite Python result: {name}!{cell.coordinate}")
                        result = NO_FINITE_BALANCE
                    expected[f"{name}!{cell.coordinate}"] = result
                else:
                    if key == C.DATE and isinstance(value, str):
                        value = date.fromisoformat(value)
                    elif key == C.MONTH and isinstance(value, str):
                        value = date.fromisoformat(value + "-01")
                    cell.value = None if value is None or pd.isna(value) else value
                if isinstance(value, (float, int)) or key in formulas:
                    cell.number_format = PERCENT_FORMAT if key in percentage_columns else DECIMAL_FORMAT
                    if key in (C.FISCAL_YEAR, C.DAYS, C.DISCOUNT_DAYS, C.NET_DAYS):
                        cell.number_format = "0"
                    elif key in (C.SHARES, C.CAP):
                        cell.number_format = "#,##0"
                elif key in (C.MONTH, C.DATE):
                    cell.number_format = "mmm-yy" if key == C.MONTH else "mm/dd/yy"
                if isinstance(value, bool):
                    cell.number_format = "General"
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="center", wrap_text=key in (C.PROVENANCE, C.DECISION))
            ws.row_dimensions[number].height = 30 if C.PROVENANCE in columns else 15
        for cell in ws[2]:
            cell.font = Font(name="Arial", size=10, bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[2].height = 32
        for key, letter in lookup.items():
            ws.column_dimensions[letter].width = TEXT_WIDTH if key in (C.BANK, C.MONTH, C.ISSUER, C.PROVENANCE, C.DECISION, C.DATE) else NUMBER_WIDTH
            if key == C.PROVENANCE:
                ws.column_dimensions[letter].width = 44
        ws.freeze_panes = "C3"
        ws.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{ws.max_row}"
        ws.sheet_view.showGridLines = False
        return ws

    fees = pd.read_csv(input_directory / "bank_fee_summary.csv").sort_values([C.BANK, C.MONTH]).to_dict("records")
    for r in fees:
        r[CARRY_COLUMN] = cfg.CARRY_ECA

    def reference(columns, key, row):
        return f"{columns[key]}{row}"

    def fee_factor(row, columns):
        return f"{columns[C.ECR]}{row}*(1-{columns[C.RESERVE]}{row})*{columns[C.DAYS]}{row}/{cfg.YEAR_DAYS}"

    fee_formulas = {
        C.COLLECTED: lambda r, c: f"={c[C.LEDGER]}{r}-{c[C.FLOAT]}{r}",
        C.ECA: lambda r, c: f"={c[C.COLLECTED]}{r}*{fee_factor(r, c)}",
        C.CARRY_IN: lambda r, c: "=0" if r == 3 else f'=IF(AND({c[CARRY_COLUMN]}{r},{c[C.BANK]}{r}={c[C.BANK]}{r-1}),{c[C.CARRY_OUT]}{r-1},0)',
        C.CARRY_OUT: lambda r, c: f"=IF({c[CARRY_COLUMN]}{r},MAX({c[C.ECA]}{r}+{c[C.CARRY_IN]}{r}-{c[C.FEES]}{r},0),0)",
        C.NET_CHARGE: lambda r, c: f"=MAX({c[C.FEES]}{r}-{c[C.ECA]}{r}-{c[C.CARRY_IN]}{r},0)",
        C.REQUIRED_BALANCE: lambda r, c: f'=IF({fee_factor(r,c)}=0,IF({c[C.FEES]}{r}=0,0,"{NO_FINITE_BALANCE}"),{c[C.FEES]}{r}/({fee_factor(r,c)}))'}
    fee_columns = [C.BANK, C.MONTH, C.LEDGER, C.FLOAT, C.COLLECTED, C.ECR, C.RESERVE, C.DAYS, C.FEES,
                   C.ECA, CARRY_COLUMN, C.CARRY_IN, C.CARRY_OUT, C.NET_CHARGE, C.REQUIRED_BALANCE,
                   C.FUND_YIELD, C.DECISION]
    sheet("Bank Fees", fees, fee_columns, fee_formulas, (C.ECR, C.RESERVE, C.FUND_YIELD))
    wc = pd.read_csv(input_directory / "working_capital.csv").to_dict("records")
    for r in wc:
        r[USE_AVERAGE_COLUMN] = r[C.BALANCE_BASIS] == "average"

    def balance_formula(key, opening_key, denominator, row, columns):
        ending = reference(columns, key, row)
        opening = reference(columns, opening_key, row)
        return f"=IF({columns[USE_AVERAGE_COLUMN]}{row},({ending}+{opening})/2,{ending})/{columns[denominator]}{row}*{columns[C.DAYS]}{row}"

    wc_formulas = {C.DSO: lambda r, c: balance_formula(C.AR, C.AR_OPEN, C.REVENUE, r, c),
                   C.DIO: lambda r, c: balance_formula(C.INVENTORY, C.INVENTORY_OPEN, C.COGS, r, c),
                   C.DPO: lambda r, c: balance_formula(C.AP, C.AP_OPEN, C.COGS, r, c),
                   C.CCC: lambda r, c: f"={c[C.DSO]}{r}+{c[C.DIO]}{r}-{c[C.DPO]}{r}"}
    sheet("CCC", wc, [C.ISSUER, C.FISCAL_YEAR, C.DAYS, C.AR_OPEN, C.AR, C.INVENTORY_OPEN,
                       C.INVENTORY, C.AP_OPEN, C.AP, C.REVENUE, C.COGS, USE_AVERAGE_COLUMN,
                       C.DSO, C.DIO, C.DPO, C.CCC, C.PROVENANCE], wc_formulas)
    discount = json.loads((input_directory / "discounts.json").read_text(encoding="utf-8"))
    sheet("Discount", [discount], [C.DISCOUNT, C.DISCOUNT_DAYS, C.NET_DAYS, C.APR, C.BORROWING],
          {C.APR: lambda r, c: f"={c[C.DISCOUNT]}{r}/(1-{c[C.DISCOUNT]}{r})*{cfg.YEAR_DAYS}/({c[C.NET_DAYS]}{r}-{c[C.DISCOUNT_DAYS]}{r})"},
          (C.DISCOUNT, C.APR, C.BORROWING))
    plan = pd.read_csv(input_directory / "buyback_plan.csv").to_dict("records")
    for r in plan:
        r[CAP_FRACTION_COLUMN] = cfg.VOLUME_FRACTION
    sheet("Buyback Cap", plan, [C.DATE, C.ADTV, CAP_FRACTION_COLUMN, C.CAP, C.SHARES, C.PRICE, C.SPEND],
          {C.CAP: lambda r, c: f"=ROUNDDOWN({c[C.ADTV]}{r}*{c[CAP_FRACTION_COLUMN]}{r},0)"}, (CAP_FRACTION_COLUMN,))
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    expected_path = path.with_name(EXPECTED_PATH.name)
    expected_path.write_text(json.dumps(expected, indent=2, allow_nan=False), encoding="utf-8")
    saved = load_workbook(path, data_only=False)
    if saved.properties.creator != cfg.AUTHOR or len(saved.sheetnames) != 4:
        raise AssertionError("Workbook metadata or sheet count differs")
    saved.close()
    return path


if __name__ == "__main__":
    print(build_workbook())
