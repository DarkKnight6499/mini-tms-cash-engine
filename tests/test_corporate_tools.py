"""Known-result fixtures, conservation checks and scoped condition boundaries."""

from copy import deepcopy
from datetime import date, timedelta
import importlib.util
from io import BytesIO
import json
from pathlib import Path
import unittest
from unittest.mock import patch

from helpers import tiny_fx

import numpy as np
from openpyxl import load_workbook
import pandas as pd

from mini_tms import config as cfg
from mini_tms.bank_fees import account_analysis, review_fees, synthetic_statements
from mini_tms.buyback import (adtv, authorization_tracker, check_safe_harbor, eps_accretion,
                              execution_plan, synthetic_blotter, volume_cap)
from mini_tms.config import C
from mini_tms.forecast import convert_flows, treasury_flows, weekly_balances
from mini_tms.models import normalize
from mini_tms.netting import run_netting, settlement_batches, synthetic_invoices
from mini_tms.payments import PaymentWorkflow, validate_pain001
from mini_tms.pooling import compare_pooling, concentrate
from mini_tms.position import daily_position
from mini_tms.rates import FXTable
from mini_tms.synthetic import generate
from mini_tms.working_capital import (annual_companyfacts, cash_conversion, collections_priority,
                                      discount_apr, discount_decisions, working_capital_table)

# CONFIG: fixed synthetic test dates and paths.
DAY = date(2025, 1, 1)
PURCHASE_DAY = date(2025, 2, 3)
WORKBOOK_TEST_DIRECTORY = cfg.ROOT / ".pytest_tmp" / "workbook_test"
FETCH_TEST_DIRECTORY = cfg.ROOT / ".pytest_tmp" / "fetch_test"


def load_script(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class NettingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = generate()
        cls.fx = FXTable(cls.data.fx)
        cls.day = cls.data.forecast_start.isoformat()
        cls.full = run_netting(synthetic_invoices(cls.data.forecast_start), cls.fx, cls.day)

    def toy(self):
        rows = [{C.INVOICE_ID: str(n), C.PAYER: p, C.RECEIVER: r, C.CURRENCY: "USD",
                 C.AMOUNT: a, C.DUE_DATE: DAY.isoformat()}
                for n, (p, r, a) in enumerate((("A", "B", 100), ("B", "C", 60), ("C", "A", 30)))]
        return run_netting(pd.DataFrame(rows), tiny_fx(), DAY.isoformat(), {e: "USD" for e in "ABC"})

    def test_toy_positions(self):
        self.assertEqual(self.toy().positions.set_index(C.ENTITY)[C.NET].to_dict(), {"A": -70, "B": 40, "C": 30})

    def test_toy_values(self):
        metrics = self.toy().metrics
        self.assertEqual(metrics["gross_value_usd"], 190)
        self.assertEqual(metrics["multilateral_value_usd"], 70)
        self.assertAlmostEqual(metrics["value_reduction_pct"] / 100, .631578947368421)

    def test_full_balance_and_count(self):
        self.assertLess(abs(self.full.positions[C.NET].sum()), 1e-6)
        self.assertLessEqual(len(self.full.instructions), len(cfg.BASE_CURRENCIES))
        self.assertEqual(len(self.full.invoices), 150)

    def test_seed_reproducibility(self):
        pd.testing.assert_frame_equal(synthetic_invoices(DAY), synthetic_invoices(DAY))

    def test_matrix_positions_reconcile(self):
        matrix = self.full.bilateral_payables
        expected = (matrix.sum(axis=0) - matrix.sum(axis=1)).rename_axis(C.ENTITY)
        np.testing.assert_allclose(expected, self.full.positions.set_index(C.ENTITY)[C.NET])

    def test_fx_only_net_foreign_amounts(self):
        self.assertTrue((self.full.fx_trades[C.CURRENCY] != "USD").all())
        self.assertEqual(len(self.full.fx_trades), (self.full.instructions[C.CURRENCY] != "USD").sum())

    def test_native_rounding_explicit(self):
        rows = self.full.instructions
        np.testing.assert_allclose(rows[C.AMOUNT], rows[C.AMOUNT].round(2))
        np.testing.assert_allclose(rows[C.SIGNED] * rows[C.FX_RATE] + rows[C.ROUNDING],
                                   rows[C.ENTITY].map(self.full.positions.set_index(C.ENTITY)[C.NET]))
        self.assertLess(rows[C.ROUNDING].abs().max(), .01)

    def test_bilateral_net_reduces_gross(self):
        metrics = self.full.metrics
        self.assertLessEqual(metrics["bilateral_value_usd"], metrics["gross_value_usd"])
        self.assertLessEqual(metrics["multilateral_value_usd"], metrics["bilateral_value_usd"])

    def test_savings_formula(self):
        m = self.full.metrics
        expected = (m["gross_payment_count"] - m["multilateral_payment_count"]) * cfg.WIRE_FEE
        expected += m["converted_volume_avoided_usd"] * cfg.FX_SPREAD_BPS / 10000
        self.assertAlmostEqual(m["estimated_savings_usd"], expected)

    def test_settlements_require_approval_and_export(self):
        workflow = PaymentWorkflow([cfg.MAKER], cfg.CHECKERS)
        for batch in settlement_batches(self.full, self.data.accounts):
            self.assertNotEqual(batch.debtor_account, batch.payments[0].creditor_account)
            workflow.create(batch)
            with self.assertRaises(PermissionError):
                workflow.export(batch.batch_id, cfg.MAKER)
            for checker in cfg.CHECKERS:
                if workflow.status(batch.batch_id) != "APPROVED":
                    workflow.approve(batch.batch_id, checker)
            xml = workflow.export(batch.batch_id, cfg.MAKER, cfg.EXPORT_TIMESTAMP)
            self.assertEqual(validate_pain001(xml, batch)["transaction_count"], 1)

    def test_duplicate_invoice_rejected(self):
        frame = self.full.invoices.copy()
        frame[C.INVOICE_ID] = "DUP"
        with self.assertRaises(ValueError):
            run_netting(frame, self.fx, self.day)


class PoolingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = generate()
        cls.fx = FXTable(cls.data.fx)
        cls.result = compare_pooling(cls.data, cls.fx)
        cls.position = daily_position(normalize(cls.data.statements, "test").statements,
                                      cls.data.accounts, [cfg.START], cls.fx)

    def test_zero_balance_and_conservation(self):
        after, _ = concentrate(self.position, self.data.accounts)
        participants = [a.account for a in self.data.accounts.values() if a.parent]
        self.assertTrue((after.set_index(C.ACCOUNT).loc[participants, C.BALANCE] == 0).all())
        np.testing.assert_allclose(after.groupby(C.CURRENCY)[C.BALANCE].sum(),
                                   self.position.groupby(C.CURRENCY)[C.BALANCE].sum(), atol=1e-9)

    def test_target_variant(self):
        after, _ = concentrate(self.position, self.data.accounts, mode="target")
        participants = [a.account for a in self.data.accounts.values() if a.parent]
        self.assertTrue((after.set_index(C.ACCOUNT).loc[participants, C.BALANCE] == cfg.POOL_TARGET).all())
        np.testing.assert_allclose(after[C.TARGET_REPORTING], after[C.TARGET] * after[C.FX_RATE])
        np.testing.assert_allclose(after[C.SURPLUS_REPORTING],
                                   after[C.REPORTING_BALANCE] - after[C.TARGET_REPORTING])

    def test_threshold_leaves_inside_band(self):
        position = self.position.copy()
        position.loc[position[C.ACCOUNT] == "A101", C.BALANCE] = 5000
        after, transfers = concentrate(position, self.data.accounts, mode="threshold")
        self.assertEqual(after.set_index(C.ACCOUNT).loc["A101", C.BALANCE], 5000)
        self.assertFalse(((transfers[C.SOURCE] == "A101") | (transfers[C.DESTINATION] == "A101")).any())

    def test_monthly_allocation_equals_group(self):
        total = self.result.comparison.set_index(C.STRUCTURE).loc["physical", C.GROUP_INTEREST]
        self.assertLess(abs(self.result.monthly_interest[C.INTEREST_USD].sum() - total), 1e-9)

    def test_ninety_days_and_three_months(self):
        self.assertEqual(self.result.balances[C.DATE].nunique(), 90)
        self.assertEqual(self.result.monthly_interest[C.MONTH].nunique(), 3)

    def test_ledger_interest_sign_and_spread(self):
        ledger = self.result.loan_ledger
        expected_rate = np.where(ledger[C.PRINCIPAL] >= 0, cfg.POOL_RATE - cfg.LOAN_SPREAD, cfg.POOL_RATE + cfg.LOAN_SPREAD)
        np.testing.assert_allclose(ledger[C.LOAN_RATE], expected_rate)
        np.testing.assert_allclose(ledger[C.INTEREST], ledger[C.PRINCIPAL] * expected_rate / 365)

    def test_running_principal_equals_transfers(self):
        balances = {}
        registry = self.data.accounts
        for r in self.result.transfers.to_dict("records"):
            source, destination = registry[r[C.SOURCE]], registry[r[C.DESTINATION]]
            if source.entity == destination.entity:
                continue
            child = r[C.SOURCE] if source.role != "concentration" else r[C.DESTINATION]
            balances[child] = balances.get(child, 0) + r[C.AMOUNT] * (1 if child == r[C.SOURCE] else -1)
        last_day = self.result.loan_ledger[C.DATE].max()
        actual = self.result.loan_ledger.loc[self.result.loan_ledger[C.DATE] == last_day].set_index(C.ACCOUNT)[C.PRINCIPAL].to_dict()
        for account, amount in balances.items():
            self.assertAlmostEqual(actual[account], amount)

    def test_notional_no_intercompany_loans(self):
        row = self.result.comparison.set_index(C.STRUCTURE).loc["notional"]
        self.assertEqual(row[C.LOANS], 0)
        self.assertGreater(row[C.GROUP_INTEREST], self.result.comparison.set_index(C.STRUCTURE).loc["standalone", C.GROUP_INTEREST])


class BankFeesTests(unittest.TestCase):
    def setUp(self):
        self.lines, self.balances, self.contract = synthetic_statements()

    def test_eca_hand_example(self):
        r = account_analysis(10000000, 500000, .02, 0, 30, 20000)
        self.assertAlmostEqual(r[C.ECA], 15616.4383561644)
        self.assertAlmostEqual(r[C.NET_CHARGE], 4383.5616438356)
        self.assertAlmostEqual(r[C.REQUIRED_BALANCE], 12166666.6666667)

    def test_no_carry_default(self):
        r = account_analysis(10000000, 0, .02, 0, 30, 100, carry_in=1000)
        self.assertEqual(r[C.NET_CHARGE], 0)
        self.assertEqual(r[C.CARRY_IN], 0)
        self.assertEqual(r[C.CARRY_OUT], 0)

    def test_carry_switch(self):
        r = account_analysis(0, 0, .02, 0, 30, 200, carry_in=1000, carry=True)
        self.assertEqual(r[C.CARRY_OUT], 800)
        self.assertEqual(r[C.NET_CHARGE], 0)

    def test_price_tolerance_boundary(self):
        frame = self.lines.head(1).copy()
        bank = frame[C.BANK].item()
        service = frame[C.SERVICE].item()
        price = self.contract.loc[(self.contract[C.BANK] == bank) & (self.contract[C.SERVICE] == service), C.CONTRACT_PRICE].item()
        for multiplier, expected in ((1.01, False), (1.0101, True), (1.0, False)):
            frame[C.UNIT_PRICE] = price * multiplier
            balances = self.balances.loc[(self.balances[C.BANK] == bank) & (self.balances[C.MONTH] == frame[C.MONTH].item())]
            _, exceptions, _, _ = review_fees(frame, balances, self.contract)
            self.assertEqual((exceptions[C.CODE] == "PRICE_ABOVE_CONTRACT").any(), expected)

    def test_deliberate_exception_types(self):
        summary, exceptions, comparison, checked = review_fees(self.lines, self.balances, self.contract)
        self.assertEqual(len(summary), 18)
        self.assertTrue({"PRICE_ABOVE_CONTRACT", "VOLUME_SPIKE", "DUPLICATE_SERVICE", "ZERO_ACTIVITY"}.issubset(set(exceptions[C.CODE])))
        self.assertEqual(len(comparison), 24)
        self.assertAlmostEqual(summary[C.FEES].sum(), checked[C.FEES].sum())

    def test_future_volume_does_not_change_prior_exceptions(self):
        _, before, _, _ = review_fees(self.lines, self.balances, self.contract)
        last_month = self.lines[C.MONTH].max()
        self.lines.loc[self.lines[C.MONTH] == last_month, C.VOLUME] *= 100
        _, after, _, _ = review_fees(self.lines, self.balances, self.contract)
        pd.testing.assert_frame_equal(before.loc[before[C.MONTH] < last_month].reset_index(drop=True),
                                      after.loc[after[C.MONTH] < last_month].reset_index(drop=True))

    def test_zero_ecr_requires_infinite_balance(self):
        r = account_analysis(100, 0, 0, 0, 30, 20)
        self.assertEqual(r[C.REQUIRED_BALANCE], float("inf"))
        self.assertEqual(r[C.NET_CHARGE], 20)

    def test_money_fund_opportunity_cost(self):
        self.assertEqual(account_analysis(100, 0, .02, 0, 30, 10)[C.DECISION], "Pay fees and invest")
        self.assertEqual(account_analysis(100, 0, .02, 0, 30, 10, fund_yield=.01)[C.DECISION], "Hold balances")


class WorkingCapitalTests(unittest.TestCase):
    def test_ending_hand_example(self):
        r = cash_conversion(100, 80, 60, 1000, 600, use_average=False)
        self.assertAlmostEqual(r[C.DSO], 36.5)
        self.assertAlmostEqual(r[C.DIO], 48.6666666667)
        self.assertAlmostEqual(r[C.DPO], 36.5)
        self.assertAlmostEqual(r[C.CCC], 48.6666666667)
        self.assertAlmostEqual(r[C.CASH_DSO], 1000 / 365)
        self.assertAlmostEqual(r[C.CASH_DIO], 600 / 365)
        self.assertAlmostEqual(r[C.CASH_DPO], 600 / 365)

    def test_average_balances(self):
        r = cash_conversion(100, 80, 60, 1000, 600, opening={C.AR: 80, C.INVENTORY: 60, C.AP: 40})
        self.assertAlmostEqual(r[C.DSO], 90 / 1000 * 365)
        self.assertAlmostEqual(r[C.DIO], 70 / 600 * 365)
        self.assertEqual(r[C.BALANCE_BASIS], "average")

    def test_average_requires_opening(self):
        with self.assertRaises(ValueError):
            cash_conversion(100, 80, 60, 1000, 600)

    def test_discount_hand_example(self):
        self.assertAlmostEqual(discount_apr(.02, 10, 30), .3724489795918367)

    def test_dynamic_and_scf_break_even(self):
        r = discount_decisions()
        self.assertTrue(r["take_discount"])
        self.assertAlmostEqual(r["dynamic_discount"], .015)
        self.assertTrue(r["scf_is_economic"])
        self.assertAlmostEqual(r["scf_net_benefit_per_invoice_dollar"], (.06 - .05) * 30 / 365)
        self.assertAlmostEqual(discount_apr(r["dynamic_break_even_discount"], 15, 30), .06)

    def test_synthetic_snapshot(self):
        result = working_capital_table(pd.read_csv(cfg.WC_FIXTURE))
        self.assertEqual(len(result), 15)
        self.assertTrue(result[C.PROVENANCE].str.contains("synthetic").all())
        self.assertTrue(np.isfinite(result[C.CCC]).all())

    def test_aging_boundaries_and_priority(self):
        days = [0, 1, 30, 31, 60, 61, 90, 91]
        ar = pd.DataFrame([{C.ENTITY: "E1", C.CURRENCY: "USD", C.INVOICE: str(age), C.AMOUNT: 100,
                            C.DUE_DATE: (DAY - timedelta(days=age)).isoformat()} for age in days])
        aging, priority = collections_priority(ar, DAY, tiny_fx())
        self.assertEqual(aging[C.USD].sum(), 800)
        self.assertEqual(priority.head(1)[C.INVOICE].item(), "91")
        self.assertEqual(priority.set_index(C.INVOICE).loc["30", C.AGING], "1_30")
        self.assertEqual(priority.set_index(C.INVOICE).loc["31", C.AGING], "31_60")


class XBRLTests(unittest.TestCase):
    def payload(self):
        facts = {}
        for metric, tags in cfg.WC_TAGS.items():
            rows = []
            for year in range(2019, 2025):
                row = {"end": f"{year}-12-31", "form": "10-K", "fp": "FY", "fy": year,
                       "filed": f"{year+1}-02-01", "accn": f"FILE-{year}", "val": 100.0}
                if metric in (C.REVENUE, C.COGS):
                    row.update(start=f"{year}-01-01", val=1000.0 if metric == C.REVENUE else 600.0)
                rows.append(row)
            facts[tags[-1]] = {"units": {"USD": rows}}
        return {"facts": {"us-gaap": facts}}

    def test_fallbacks_and_five_years(self):
        table = annual_companyfacts(self.payload(), "TEST")
        self.assertEqual(len(table), 5)
        self.assertEqual(table[C.FISCAL_YEAR].min(), 2020)
        self.assertTrue((table["revenue_tag"] == "Revenues").all())
        self.assertTrue((table["cogs_tag"] == "CostOfGoodsSold").all())

    def test_reject_quarter_and_wrong_form(self):
        payload = self.payload()
        rows = payload["facts"]["us-gaap"]["Revenues"]["units"]["USD"]
        rows.extend([{**rows[-1], "start": "2024-10-01", "val": 1, "filed": "2026-01-01"},
                     {**rows[-1], "form": "10-Q", "val": 2, "filed": "2027-01-01"}])
        self.assertEqual(annual_companyfacts(payload, "TEST").tail(1)[C.REVENUE].item(), 1000)

    def test_comparative_fiscal_year_not_used(self):
        payload = self.payload()
        rows = payload["facts"]["us-gaap"]["Revenues"]["units"]["USD"]
        rows.append({**rows[-1], "fy": 2025, "filed": "2026-02-01", "val": 1100})
        result = annual_companyfacts(payload, "TEST").tail(1)
        self.assertEqual(result[C.FISCAL_YEAR].item(), 2024)
        self.assertEqual(result[C.REVENUE].item(), 1100)

    def test_missing_opening_not_fabricated(self):
        payload = self.payload()
        del payload["facts"]["us-gaap"]["InventoryNet"]["units"]["USD"][0]
        with self.assertRaisesRegex(ValueError, "missing opening"):
            annual_companyfacts(payload, "TEST")

    def test_misaligned_annual_cogs_rejected(self):
        payload = self.payload()
        payload["facts"]["us-gaap"]["CostOfGoodsSold"]["units"]["USD"][-1]["start"] = "2024-01-02"
        with self.assertRaisesRegex(ValueError, "missing aligned annual cogs"):
            annual_companyfacts(payload, "TEST")

    def test_fallback_after_period_alignment(self):
        payload = self.payload()
        rows = deepcopy(payload["facts"]["us-gaap"]["CostOfGoodsSold"]["units"]["USD"])
        for row in rows:
            row["start"] = row["start"].replace("01-01", "01-02")
        payload["facts"]["us-gaap"]["CostOfGoodsAndServicesSold"] = {"units": {"USD": rows}}
        table = annual_companyfacts(payload, "TEST")
        self.assertTrue((table["cogs_tag"] == "CostOfGoodsSold").all())

    def test_fiscal_year_anchor_across_fallback_tags(self):
        payload = self.payload()
        latest = deepcopy(payload["facts"]["us-gaap"]["Revenues"]["units"]["USD"][-1])
        latest.update(fy=2025, filed="2026-02-01", val=1200)
        payload["facts"]["us-gaap"]["RevenueFromContractWithCustomerExcludingAssessedTax"] = {"units": {"USD": [latest]}}
        table = annual_companyfacts(payload, "TEST").tail(1)
        self.assertEqual(table[C.FISCAL_YEAR].item(), 2024)
        self.assertEqual(table[C.REVENUE].item(), 1200)

    @unittest.skipUnless(cfg.SNAPSHOT.exists(), "SEC snapshot absent")
    def test_optional_real_snapshot(self):
        table = working_capital_table(pd.read_csv(cfg.SNAPSHOT))
        self.assertEqual(set(table[C.ISSUER]), set(cfg.WC_ISSUERS))
        self.assertEqual(len(table), 15)


class FetchPublicationTests(unittest.TestCase):
    def test_incomplete_fetch_preserves_published_snapshot(self):
        downloader = load_script("treasury_fetcher", cfg.ROOT / "examples" / "fetch_working_capital.py")
        folder = FETCH_TEST_DIRECTORY
        folder.mkdir(parents=True, exist_ok=True)
        snapshot = folder / cfg.SNAPSHOT.name
        snapshot.write_text("existing complete snapshot", encoding="utf-8")
        tickers = {str(n): {"ticker": issuer, "cik_str": n + 1}
                   for n, issuer in enumerate(cfg.WC_ISSUERS)}
        raw = json.dumps(XBRLTests().payload()).encode("utf-8")
        responses = [BytesIO(json.dumps(tickers).encode("utf-8")), BytesIO(raw),
                     OSError("Synthetic offline failure"), BytesIO(raw)]
        with (patch.object(cfg, "DATA_DIRECTORY", folder), patch.object(cfg, "SNAPSHOT", snapshot),
              patch.object(downloader, "MANIFEST", folder / "manifest.json"),
              patch.object(downloader, "urlopen", side_effect=responses),
              patch.object(downloader.time, "sleep"), patch("builtins.print")):
            self.assertFalse(downloader.fetch_working_capital())
        self.assertEqual(snapshot.read_text(encoding="utf-8"), "existing complete snapshot")
        manifest = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
        self.assertNotIn("snapshot", manifest)
        self.assertEqual([r["status"] for r in manifest["issuers"]], ["ok", "failed", "ok"])
        self.assertEqual(len(manifest["requests"]), 4)
        for request in manifest["requests"]:
            if request["status"] == "ok":
                self.assertEqual(len(request["sha256"]), 64)
                self.assertTrue((folder / request["path"]).is_file())


class BuybackTests(unittest.TestCase):
    def setUp(self):
        self.blotter, self.history = synthetic_blotter(PURCHASE_DAY)

    def trade(self, **updates):
        row = {C.TRADE_ID: "TEST", C.DATE: PURCHASE_DAY.isoformat(), C.TIME: "10:00",
               C.BROKER: "B1", C.SHARES: 250000, C.PRICE: 50, C.BID: 50, C.LAST: 49,
               C.OPENING: False, C.BLOCK: False, C.BLOCK_QUALIFIED: False}
        row.update(updates)
        return row

    def check(self, rows, **kwargs):
        return check_safe_harbor(pd.DataFrame(rows), self.history, **kwargs)

    def test_adtv_and_daily_cap(self):
        self.assertEqual(adtv(self.history, PURCHASE_DAY), 1000000)
        self.assertEqual(volume_cap(self.history, PURCHASE_DAY), 250000)

    def test_260000_flagged_cap_boundary_allowed(self):
        self.assertTrue(self.check([self.trade()]).empty)
        self.assertIn("(b)(4)", set(self.check([self.trade(**{C.SHARES: 260000})])[C.RULE]))

    def test_opening_and_large_late_purchase(self):
        self.assertIn("(b)(2)(i)", set(self.check([self.trade(**{C.OPENING: True})])[C.RULE]))
        self.assertIn("(b)(2)(ii)", set(self.check([self.trade(**{C.TIME: "15:55"})])[C.RULE]))

    def test_small_security_thirty_minutes(self):
        result = self.check([self.trade(**{C.TIME: "15:35"})], public_float=149999999)
        self.assertIn("(b)(2)(iii)", set(result[C.RULE]))

    def test_blackout_boundary(self):
        self.assertTrue(self.check([self.trade(**{C.TIME: "15:49"})]).empty)
        self.assertIn("(b)(2)(ii)", set(self.check([self.trade(**{C.TIME: "15:50"})])[C.RULE]))

    def test_two_brokers(self):
        rows = [self.trade(**{C.SHARES: 100}), self.trade(**{C.TRADE_ID: "SECOND", C.BROKER: "B2", C.SHARES: 100})]
        self.assertIn("(b)(1)", set(self.check(rows)[C.RULE]))

    def test_price_independent_cap(self):
        self.assertIn("(b)(3)", set(self.check([self.trade(**{C.PRICE: 50.01})])[C.RULE]))
        self.assertTrue(self.check([self.trade(**{C.PRICE: 51, C.LAST: 51})]).empty)

    def test_weekly_block_exception(self):
        row = self.trade(**{C.SHARES: 500000, C.BLOCK: True, C.BLOCK_QUALIFIED: True})
        self.assertTrue(self.check([row]).empty)
        second = {**row, C.DATE: (PURCHASE_DAY + timedelta(days=1)).isoformat(), C.TRADE_ID: "SECOND"}
        self.assertIn("(b)(4)", set(self.check([row, second])[C.RULE]))

    def test_block_requires_one_purchase_and_verification(self):
        row = self.trade(**{C.SHARES: 500000, C.BLOCK: True})
        self.assertIn("(b)(4)", set(self.check([row])[C.RULE]))
        row[C.BLOCK_QUALIFIED] = True
        other = self.trade(**{C.TRADE_ID: "SECOND", C.SHARES: 1})
        self.assertIn("(b)(4)", set(self.check([row, other])[C.RULE]))

    def test_prior_block_in_same_week(self):
        row = self.trade(**{C.SHARES: 500000, C.BLOCK: True, C.BLOCK_QUALIFIED: True})
        self.assertIn("(b)(4)", set(self.check([row], prior_block_dates=(PURCHASE_DAY,))[C.RULE]))

    def test_adtv_excludes_current_week_and_blocks(self):
        self.history.loc[self.history[C.DATE] == PURCHASE_DAY.isoformat(), C.VOLUME] = 900000000
        self.assertEqual(adtv(self.history, PURCHASE_DAY), 1000000)
        prior = self.history[C.DATE] < PURCHASE_DAY.isoformat()
        self.history.loc[prior, C.VOLUME_EXCLUDED] = 100000
        self.assertEqual(adtv(self.history, PURCHASE_DAY), 900000)

    def test_incomplete_adtv_history_rejected(self):
        with self.assertRaisesRegex(ValueError, "four-calendar-week"):
            adtv(self.history.head(5), PURCHASE_DAY)

    def test_eps_hand_example(self):
        r = eps_accretion(1000, 100, 500, 50, .04, .25)
        self.assertEqual(r["pro_forma_shares"], 90)
        self.assertEqual(r["pro_forma_net_income"], 985)
        self.assertAlmostEqual(r["pro_forma_eps"], 10.9444444444)
        self.assertEqual(r["base_eps"], 10)

    def test_authorization(self):
        r = authorization_tracker(pd.DataFrame([self.trade()]))
        self.assertEqual(r["cumulative_spend"], 12500000)
        self.assertEqual(r["remaining_capacity"], cfg.AUTHORIZATION - 12500000)
        self.assertEqual(r["average_price"], 50)

    def test_plan_honors_capacity(self):
        plan = execution_plan([PURCHASE_DAY], self.history, spend=20000000, remaining=10000000)
        self.assertEqual(plan[C.SPEND].sum(), 10000000)
        self.assertTrue((plan[C.SHARES] <= plan[C.CAP]).all())


class ForecastHookTests(unittest.TestCase):
    def test_settlement_center_legs_cancel_and_buyback_outflow(self):
        data = generate()
        fx = FXTable(data.fx)
        netting = run_netting(synthetic_invoices(data.forecast_start), fx, data.forecast_start.isoformat())
        blotter, history = synthetic_blotter(data.forecast_start)
        plan = execution_plan([data.forecast_start], history)
        flows = treasury_flows(netting.instructions, plan, data.forecast_start)
        converted = convert_flows(flows, fx, data.forecast_start.isoformat())
        self.assertAlmostEqual(converted.loc[converted[C.LINE] == "netting", C.AMOUNT_REPORTING].sum(), 0)
        self.assertEqual(converted.loc[converted[C.LINE] == "buyback", C.AMOUNT_REPORTING].sum(), -500000)
        weekly = weekly_balances(converted, {e: 0 for e in cfg.BASE_CURRENCIES}, {e: 0 for e in cfg.BASE_CURRENCIES})
        self.assertAlmostEqual(weekly.loc[weekly[C.WEEK] == 12, C.CLOSING].sum(), -500000)


class WorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.builder = load_script("treasury_builder", cfg.ROOT / "excel" / "build_workbook.py")
        cls.recalculator = load_script("treasury_recalculator", cfg.ROOT / "excel" / "reconcile_workbook.py")

    def test_live_formulas_and_plain_metadata(self):
        source = WORKBOOK_TEST_DIRECTORY / "inputs"
        source.mkdir(parents=True, exist_ok=True)
        lines, balances, contract = synthetic_statements()
        fees, _, _, _ = review_fees(lines, balances, contract)
        fees.to_csv(source / "bank_fee_summary.csv", index=False)
        working_capital_table(pd.read_csv(cfg.WC_FIXTURE)).to_csv(source / "working_capital.csv", index=False)
        (source / "discounts.json").write_text(json.dumps(discount_decisions()), encoding="utf-8")
        _, history = synthetic_blotter(PURCHASE_DAY)
        execution_plan([PURCHASE_DAY], history).to_csv(source / "buyback_plan.csv", index=False)
        path = self.builder.build_workbook(source, WORKBOOK_TEST_DIRECTORY / "tools.xlsx")
        book = load_workbook(path, data_only=False)
        self.assertEqual(book.properties.creator, cfg.AUTHOR)
        self.assertEqual(set(book.sheetnames), {"Bank Fees", "CCC", "Discount", "Buyback Cap"})
        expected = json.loads(path.with_name("reconciliation_expected.json").read_text())
        for address in expected:
            name, cell = address.split("!")
            self.assertEqual(book[name][cell].data_type, "f")
        for sheet in book:
            self.assertIsNone(sheet.sheet_properties.tabColor)
            for row in sheet:
                for cell in row:
                    self.assertIsNone(cell.fill.patternType)
                    if cell.font.color is not None:
                        self.assertNotEqual(cell.font.color.type, "rgb")
        book.close()

    def test_libreoffice_reconciliation(self):
        if self.recalculator.libreoffice_path() is None:
            self.skipTest("LibreOffice absent")
        self.test_live_formulas_and_plain_metadata()
        path = WORKBOOK_TEST_DIRECTORY / "tools.xlsx"
        self.assertEqual(self.recalculator.reconcile(path, WORKBOOK_TEST_DIRECTORY / "recalculated")["status"], "passed")

    def test_changed_ecr_recalculates(self):
        if self.recalculator.libreoffice_path() is None:
            self.skipTest("LibreOffice absent")
        self.test_live_formulas_and_plain_metadata()
        path = WORKBOOK_TEST_DIRECTORY / "tools.xlsx"
        book = load_workbook(path)
        sheet = book["Bank Fees"]
        columns = {cell.value: cell.column_letter for cell in sheet[2]}
        inputs = {key: sheet[f"{letter}3"].value for key, letter in columns.items()}
        sheet[f"{columns[C.ECR]}3"] = inputs[C.ECR] * 2
        book.save(path)
        book.close()
        expected_path = path.with_name("reconciliation_expected.json")
        expected = json.loads(expected_path.read_text())
        updated = account_analysis(inputs[C.LEDGER], inputs[C.FLOAT], inputs[C.ECR] * 2,
                                    inputs[C.RESERVE], inputs[C.DAYS], inputs[C.FEES])
        for key in (C.ECA, C.NET_CHARGE, C.REQUIRED_BALANCE):
            expected[f"Bank Fees!{columns[key]}3"] = updated[key]
        expected_path.write_text(json.dumps(expected), encoding="utf-8")
        self.assertEqual(self.recalculator.reconcile(path, WORKBOOK_TEST_DIRECTORY / "changed_ecr")["status"], "passed")
        book = load_workbook(path, data_only=True)
        self.assertAlmostEqual(book["Bank Fees"][f"{columns[C.ECA]}3"].value, updated[C.ECA])
        book.close()

    def test_zero_ecr_and_zero_fees_recalculate(self):
        if self.recalculator.libreoffice_path() is None:
            self.skipTest("LibreOffice absent")
        self.test_live_formulas_and_plain_metadata()
        source = WORKBOOK_TEST_DIRECTORY / "inputs"
        lines, balances, contract = synthetic_statements()
        balances[C.ECR] = 0.0
        lines.loc[lines[C.MONTH] == lines[C.MONTH].min(), C.UNIT_PRICE] = 0.0
        fees, _, _, _ = review_fees(lines, balances, contract)
        fees.to_csv(source / "bank_fee_summary.csv", index=False)
        path = self.builder.build_workbook(source, WORKBOOK_TEST_DIRECTORY / "zero_ecr.xlsx")
        self.assertEqual(self.recalculator.reconcile(path, WORKBOOK_TEST_DIRECTORY / "zero_ecr")["status"], "passed")
        book = load_workbook(path, data_only=True)
        sheet = book["Bank Fees"]
        columns = {cell.value: cell.column_letter for cell in sheet[2]}
        self.assertEqual(sheet[f"{columns[C.REQUIRED_BALANCE]}3"].value, 0)
        self.assertEqual(sheet[f"{columns[C.REQUIRED_BALANCE]}4"].value, self.builder.NO_FINITE_BALANCE)
        self.assertEqual(sheet[f"{columns[C.ECA]}4"].value, 0)
        book.close()


if __name__ == "__main__":
    unittest.main()
